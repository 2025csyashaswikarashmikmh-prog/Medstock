"""PDF (reportlab) and Excel (openpyxl) generation: invoices, lists, exports."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import db

TEAL = colors.HexColor("#0f766e")
LIGHT = colors.HexColor("#f0fdfa")
styles = getSampleStyleSheet()
small = ParagraphStyle("small", parent=styles["Normal"], fontSize=8.5, leading=11)
cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8.5, leading=10)


def _p(text, style=small):
    # Paragraph parses mini-HTML, so escape user-entered text
    text = str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return Paragraph(text, style)


def _money(n):
    return f"{n:,.2f}"          # "Rs." is written in labels: the default PDF font has no rupee glyph


def _grid(data, widths, header_rows=1, align_right_from=None):
    t = Table(data, colWidths=widths, repeatRows=header_rows)
    style = [
        ("BACKGROUND", (0, 0), (-1, header_rows - 1), TEAL),
        ("TEXTCOLOR", (0, 0), (-1, header_rows - 1), colors.white),
        ("FONTNAME", (0, 0), (-1, header_rows - 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#a7d7cc")),
        ("ROWBACKGROUNDS", (0, header_rows), (-1, -1), [colors.white, LIGHT]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    if align_right_from is not None:
        style.append(("ALIGN", (align_right_from, 0), (-1, -1), "RIGHT"))
    t.setStyle(TableStyle(style))
    return t


def invoice_pdf(bill):
    shop = {k: db.get_setting(k) for k in
            ("shop_name", "shop_address", "shop_phone", "shop_gstin", "shop_dl_no", "invoice_footer")}
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=f"Invoice {bill['bill_no']}")
    story = [
        Paragraph(f"<b>{shop['shop_name']}</b>".replace("&", "&amp;"),
                  ParagraphStyle("t", parent=styles["Title"], textColor=TEAL, fontSize=20, spaceAfter=2)),
        _p(shop["shop_address"]),
        _p(" | ".join(x for x in (
            f"Phone: {shop['shop_phone']}" if shop["shop_phone"] else "",
            f"GSTIN: {shop['shop_gstin']}" if shop["shop_gstin"] else "",
            f"Drug Licence: {shop['shop_dl_no']}" if shop["shop_dl_no"] else "") if x)),
        Spacer(1, 6),
        Paragraph("<b>TAX INVOICE</b>", ParagraphStyle("h", parent=styles["Heading3"], alignment=1)),
    ]
    meta = [
        [_p(f"Invoice No: {bill['bill_no']}"), _p(f"Date: {bill['created_at'].replace('T', ' ')}")],
        [_p(f"Customer: {bill['customer_name'] or 'Walk-in'}" + (f"  ({bill['customer_phone']})" if bill['customer_phone'] else "")),
         _p(f"Doctor: {bill['doctor_name'] or '-'}    Payment: {bill['payment_mode']}")],
    ]
    mt = Table(meta, colWidths=[90 * mm, 92 * mm])
    mt.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, TEAL), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [mt, Spacer(1, 8)]

    head = ["#", "Medicine", "Batch", "Exp", "Qty", "MRP", "GST%", "Amount"]
    rows = [head]
    for n, i in enumerate(bill["items"], 1):
        rows.append([n, _p(i["name"], cell), i["batch"], i["expiry"][:7], i["qty"], _money(i["price"]),
                     f"{i['gst_rate']:g}", _money(i["total"])])
    story.append(_grid(rows, [8 * mm, 62 * mm, 24 * mm, 16 * mm, 12 * mm, 18 * mm, 14 * mm, 28 * mm]))
    story.append(Spacer(1, 8))

    # GST summary by rate (needed on a proper tax invoice)
    by_rate = {}
    for i in bill["items"]:
        d = by_rate.setdefault(i["gst_rate"], [0.0, 0.0])
        d[0] += i["taxable"]
        d[1] += i["gst_amount"]
    gst_rows = [["GST %", "Taxable value", "CGST", "SGST"]]
    for rate, (tx, g) in sorted(by_rate.items()):
        gst_rows.append([f"{rate:g}%", _money(tx), _money(g / 2), _money(g / 2)])
    left = _grid(gst_rows, [16 * mm, 30 * mm, 22 * mm, 22 * mm])

    disc = f" (after {bill['discount_pct']:g}% discount)" if bill["discount_pct"] else ""
    tot = [
        ["Taxable value", _money(bill["taxable"])],
        ["CGST", _money(bill["cgst"])],
        ["SGST", _money(bill["sgst"])],
        ["Round off", _money(bill["round_off"])],
        [_p(f"GRAND TOTAL (Rs.){disc}", cell), _money(bill["total"])],
    ]
    right = Table(tot, colWidths=[50 * mm, 28 * mm])
    right.setStyle(TableStyle([
        ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("LINEABOVE", (0, -1), (-1, -1), 0.8, TEAL),
        ("BACKGROUND", (0, -1), (-1, -1), LIGHT),
    ]))
    side = Table([[left, right]], colWidths=[96 * mm, 86 * mm])
    side.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [side, Spacer(1, 10), _p("Prices are MRP inclusive of GST."), _p(shop["invoice_footer"])]
    doc.build(story)
    return buf.getvalue()


def table_pdf(title, headers, rows, subtitle="", landscape_mode=True, widths=None):
    buf = BytesIO()
    size = landscape(A4) if landscape_mode else A4
    doc = SimpleDocTemplate(buf, pagesize=size, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    story = [Paragraph(f"<b>{db.get_setting('shop_name')}</b> - {title}".replace("&", "&amp;"), styles["Heading2"])]
    if subtitle:
        story.append(_p(subtitle))
    story.append(Spacer(1, 6))
    data = [headers] + [[_p(c, cell) if isinstance(c, str) else c for c in r] for r in rows]
    if len(data) == 1:
        story.append(_p("Nothing to show."))
    else:
        story.append(_grid(data, widths))
    doc.build(story)
    return buf.getvalue()


def xlsx_bytes(sheet, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F766E")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):          # user text such as "=cmd" must stay text, never a formula
        for c in row:
            if isinstance(c.value, str) and c.value[:1] in "=+-@":
                c.data_type = "s"
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 45)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
